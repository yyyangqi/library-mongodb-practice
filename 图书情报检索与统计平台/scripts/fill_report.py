# -*- coding: utf-8 -*-
"""填写实验报告模板（模板原生栏目）：顶部下划线填空 + 表格单元格内容。
删除单元格空占位段落，压缩大最小行高以控制页数。
用法：python fill_report.py [输出路径]（默认写到数据2403班-202414870302-杨淇.docx）"""
import sys
from docx import Document
from docx.oxml.ns import qn

DEFAULT = r"C:\Users\26295\Doubao\chats\2026-10-21\new-chat\图书情报检索与统计平台\数据2403班-202414870302-杨淇.docx"
P = sys.argv[1] if len(sys.argv) > 1 else DEFAULT
doc = Document(P)

FONT_MAP = {'KaiTi_GB2312': '楷体', 'FangSong_GB2312': '仿宋', 'SimHei': '黑体', 'SimSun': '宋体'}

def fix_fonts(doc):
    """将模板遗留的旧版字体名替换为现代 Word 可渲染的规范字体。"""
    def fix_run(r):
        rPr = r._element.rPr
        if rPr is None:
            return
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            return
        for k, v in list(rFonts.attrib.items()):
            if k.endswith('eastAsia') and v in FONT_MAP:
                rFonts.set(k, FONT_MAP[v])
    for p in doc.paragraphs:
        for r in p.runs:
            fix_run(r)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        fix_run(r)

def fill_underline(para, groups):
    runs = para.runs
    for start, end, value in groups:
        if start < len(runs):
            runs[start].text = value
            for i in range(start + 1, min(end, len(runs) - 1) + 1):
                runs[i].text = ''
        else:
            runs.append(value)

p1 = doc.paragraphs[1]
fill_underline(p1, [(3, 8, '数据2403班'), (10, 14, '202414870302'), (16, 20, '杨淇')])
p2 = doc.paragraphs[2]
fill_underline(p2, [(1, 1, '2026年10月21日'), (5, 5, '1')])

tbl = doc.tables[0]
rows = tbl.rows

def set_cell(cell, lines):
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)
    for ln in lines:
        cell.add_paragraph(ln)

def shrink_row_height(row):
    trPr = row._tr.find(qn('w:trPr'))
    if trPr is None:
        return
    for child in list(trPr):
        if child.tag.endswith('}trHeight'):
            trPr.remove(child)
    if len(trPr) == 0:
        row._tr.remove(trPr)

for ri in (2, 3, 4, 5, 6):
    shrink_row_height(rows[ri])

set_cell(rows[0].cells[1], ['NoSQL数据库技术与应用'])
set_cell(rows[1].cells[1], ['图书情报检索与统计平台（基于MongoDB）'])

set_cell(rows[2].cells[1], [
    '实验目的与要求',
    '1. 掌握 MongoDB 的安装、部署与服务启动方法。',
    '2. 掌握集合与特殊集合（时间序列、固定集合、视图）的创建。',
    '3. 掌握 CRUD、条件检索、投影与排序。',
    '4. 掌握聚合管道、索引与权限管理。',
    '5. 掌握基于 PyMongo 的应用开发。',
    '6. 掌握 Compass 图形化验证与数据导出。',
])

set_cell(rows[3].cells[1], [
    '实验环境的配置',
    'MongoDB 8.0.4（ZIP，解压至 D:\\MongoDB）、mongosh 2.2.10、Python 3.13 + PyMongo 4.19.0。',
    '数据目录 D:\\MongoDB\\data\\db，监听 127.0.0.1:27017；Windows 平台（PowerShell + mongosh + Python）。',
])

set_cell(rows[4].cells[1], [
    '实验内容与实验步骤和结果',
    '【任务一 环境部署】安装 MongoDB 8.0.4 与 mongosh，配置 PATH，启动 mongod，mongosh 连接成功。',
    '【任务二 数据建模】创建 library 库与 books(20)/readers(16)/borrow_logs(66)/category_stats(5) 四集合及三种特殊集合，导入数据；borrow_logs 时间序列集合保留期设为 5 年（157680000 秒），避免早期借阅记录被 TTL 自动清理。',
    '【任务三 数据检索】实现 Q1—Q16 检索、U1—U4 更新（$inc/$mul/$currentDate/$unset）、D1—D3 删除，全部成功。典型结果：Q1=7、Q2=8、Q3=1、Q5=16、Q14 评分前 5。',
    '【任务四 统计、索引与权限】聚合 A1—A9（分类统计、出版社前5、借阅年份趋势、$lookup 学院排行、$merge 落库）、索引 I1—I5 与 explain 验证、权限 libReader/libWriter 验证。',
    '【任务五 PyMongo 应用】开发 6 个 Python 源文件，verify_app 验证 9 项功能通过，main.py 运行正常。',
    '【任务六 验证与导出】mongosh 导出 json（books 20 条、category_stats 5 条），getCollectionInfos 查看集合类型标记（timeseries/capped/view）。',
])

set_cell(rows[5].cells[1], [
    '（关键命令示例详见 code 与 scripts 目录；详细实现见技术文档。）',
])

set_cell(rows[6].cells[1], [
    '心得体会',
    '通过本实训系统掌握 MongoDB 全流程，体会到文档型数据库的灵活性、特殊集合的应用价值、索引对查询性能的关键作用，以及权限管理对数据安全的保障。',
])

fix_fonts(doc)
doc.save(P)
print("已按实验报告模板原生栏目填写并保存：", P)
