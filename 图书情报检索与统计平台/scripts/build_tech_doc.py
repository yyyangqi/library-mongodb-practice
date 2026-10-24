# -*- coding: utf-8 -*-
"""生成技术文档（单独成文）：图书情报检索与统计平台，四章节，3-4页。
A4、2.5cm 边距；标题黑体18pt；四章节 Heading 黑体14pt；正文宋体12pt、1.5倍行距、首行缩进。"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = r"C:\Users\26295\Doubao\chats\2026-10-21\new-chat\图书情报检索与统计平台\数据2403班-202414870302-杨淇-技术文档.docx"
doc = Document()

# 页面设置 A4, 2.5cm 边距
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2.5)
sec.left_margin = sec.right_margin = Cm(2.5)

def set_font(run, chinese, size, bold=False):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), chinese)
    run.font.size = Pt(size)
    run.font.bold = bold

def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, '黑体', 18, bold=True)
    return p

def add_info(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(text)
    set_font(r, '楷体', 12)
    return p

def add_heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, '黑体', 14, bold=True)
    return p

def add_sub(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_font(r, '黑体', 12, bold=True)
    return p

def add_para(text, indent=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(3)
    if indent:
        pf.first_line_indent = Pt(24)
    r = p.add_run(text)
    set_font(r, '宋体', 12)
    return p

add_title('图书情报检索与统计平台\n技术文档')
add_info('班级：数据2403班　　学号：202414870302　　姓名：杨淇　　日期：2026年10月21日')

add_heading('一、项目概述')
add_para('项目背景：本项目基于 MongoDB 构建《图书情报检索与统计平台》，围绕图书、读者、借阅与统计业务，实现数据的管理、检索与统计分析，验证文档型 NoSQL 数据库在图书情报场景中的实际应用价值。')
add_para('业务目标：完成 MongoDB 部署、数据建模、数据检索、统计与权限、PyMongo 应用、Compass 验证六项任务，覆盖 16 项业务功能（F1—F16），最终以命令行与图形化双通道核对业务数据。')
add_para('交付物：数据文件（books/readers/borrow_logs 数据集及导出 json）、6 个可运行的 Python 源文件、本技术文档（四章节）及覆盖六任务的截图。')

add_heading('二、环境与数据建模')
add_sub('（一）部署关键参数')
add_para('MongoDB 8.0.4（ZIP 方式解压至 D:\\MongoDB）、mongosh 2.2.10、Python 3.13 + PyMongo 4.19.0。数据目录 D:\\MongoDB\\data\\db，监听 127.0.0.1:27017，以无认证模式运行；权限实验时以 --auth 重启验证后恢复无认证模式。')
add_sub('（二）四个业务集合')
add_para('books（20 条）、readers（16 条）、borrow_logs（66 条，时间序列集合）、category_stats（聚合落库 5 条）。')
add_sub('（三）三种特殊集合')
add_para('时间序列集合 borrow_logs：timeField=borrow_date，granularity=hours，expireAfterSeconds=157680000（保留期 5 年）；固定集合 book_logs：capped，1MB/1000 条；视图 book_in_stock：过滤 stock>0，共 19 条。')
add_sub('（四）索引设计')
add_para('isbn 唯一索引、title/category 单键索引、tags 多键索引、category+price 复合索引共 5 项；通过 explain 对比，无索引（author）扫描 20 条 vs 有索引（isbn）扫描 1 条，验证索引对查询性能的提升。')

add_heading('三、功能实现')
add_sub('【任务一 环境部署】')
add_para('下载安装 MongoDB 8.0.4 与 mongosh，配置 PATH，启动 mongod（日志出现 "mongod startup complete"），mongosh 连接成功。')
add_sub('【任务二 数据建模】')
add_para('创建 library 库与 books/readers/borrow_logs/category_stats 四集合及三种特殊集合，导入数据（books=20、readers=16、borrow_logs=66）。数据建模：borrow_logs 为时间序列集合，保留期从当前时间往前倒推计算；原 2 年保留期会自动清理 2024 年 10 月以前的借阅记录（由 66 条降至 26 条），按要求将保留期由 2 年改为 5 年（expireAfterSeconds=157680000 秒），重建 borrow_logs 集合后重新导入，保留全部 66 条。')
add_sub('【任务三 数据检索】')
add_para('实现 Q1—Q16 条件检索、U1—U4 更新（$inc/$mul/$currentDate/$unset）、D1—D3 删除（deleteOne/deleteMany/replaceOne），全部成功。典型结果：Q1=7、Q2=8、Q3=1、Q5=16、Q7=7、Q11=12、Q12=17；Q14 评分前 5 为《活着》《哥德尔证明》《深入理解计算机系统》《平凡的世界》《费马大定理》。')
add_sub('【任务四 统计、索引与权限】')
add_para('聚合 A1—A9：分类统计（计算机 7/文学 5/数学 3/经济管理 3/外语 2）、出版社前 5、销售额、借阅年份趋势（2023=20/2024=28/2025=18）、读者借阅前 5、$lookup 学院排行、$merge 落库 category_stats=5。索引 I1—I5 创建并 explain 验证，唯一索引拦截重复 ISBN（E11000）。权限：无认证读被拒、libReader 读=20 写被拒、libWriter 读写成功。')
add_sub('【任务五 PyMongo 应用】')
add_para('开发 6 个 Python 源文件（连接封装/图书/读者/借阅/统计/菜单主程序 main.py 含 13 项功能）；verify_app.py 对 9 项功能验证通过，main.py 运行正常。')
add_sub('【任务六 图形化验证与数据导出】')
add_para('mongosh 导出 books_export.json（20 条）、category_stats_export.json（5 条）；db.getCollectionInfos() 查看集合类型标记（borrow_logs=timeseries、book_logs=capped、book_in_stock=view）；命令行与验证结果（Q1/Q2/Q7/Q14/A1/视图）核对一致，亦可在 Compass 图形化界面等价复查。')

add_heading('四、问题与核对')
add_sub('（一）调试难点与排查')
add_para('1. mongosh 脚本中 use library 非法，改用 db.getSiblingDB("library")；cat() 不存在，改用 require("fs").readFileSync 读取导入文件。')
add_para('2. 时间序列集合 TTL 保留期从当前时间往前倒推，原 2 年保留期会自动清理 2024 年 10 月以前的借阅记录（66→26）；将 expireAfterSeconds 由 2 年调整为 5 年（157680000 秒），重建集合后重新导入，保留全部 66 条。')
add_para('3. 索引优化：explain 对比无索引（author）扫 20 条 vs 有索引（isbn）扫 1 条，验证索引显著提升查询性能；唯一索引成功拦截重复 ISBN（E11000）。')
add_para('4. 权限验证需以 --auth 重启 mongod，认证连接须指定 authSource=admin；验证后恢复无认证模式。')
add_sub('（二）命令行与 Compass 核对')
add_para('Q1=7、Q2=8、Q7=7、Q14 前 5、A1 五组、视图 book_in_stock=19，与命令行及 Compass 验证结果一致。')
add_para('总结：通过本实训系统掌握 MongoDB 全流程，深刻体会文档型数据库的灵活性、特殊集合的应用价值、索引对性能的关键作用及权限管理对数据安全的保障。')

doc.save(OUT)
print("技术文档已生成：", OUT)
